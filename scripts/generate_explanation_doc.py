from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "documento-explicativo-proyecto.docx"

BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
MUTED = RGBColor(89, 89, 89)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 14, 7),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 12, DARK, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_title_block(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Documento explicativo del proyecto final")
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(32, 35, 31)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run("Generador de Recetas con Inventario | Tecnologias Web")
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.color.rgb = BLUE

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    run = meta.add_run(
        "Integrantes: Sean Paul Marquez Toro, Reyner David Barbosa de la Rosa y Ruben Andres Corro Blanco"
    )
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.color.rgb = MUTED


def add_key_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(4.8)

    headers = table.rows[0].cells
    set_cell_text(headers[0], "Entregable", True)
    set_cell_text(headers[1], "Resultado final", True)
    set_cell_shading(headers[0], "F2F4F7")
    set_cell_shading(headers[1], "F2F4F7")

    rows = [
        ("Repositorio", "https://github.com/rubencorrob-sudo/generador-recetas-inventario"),
        ("Aplicacion", "https://www.recetasruben.xyz"),
        ("Swagger", "https://www.recetasruben.xyz/docs"),
        ("PDF tecnico", "docs/informe-proyecto-recetas.pdf"),
        ("Respaldo temporal", "https://54-236-36-56.sslip.io"),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], label, True)
        cells[1].text = ""
        p = cells[1].paragraphs[0]
        if value.startswith("https://"):
            add_hyperlink(p, value, value)
        else:
            p.add_run(value)
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_bullets(doc: Document, items) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbered(doc: Document, items) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def build() -> None:
    doc = Document()
    style_document(doc)
    add_title_block(doc)

    doc.add_heading("Resumen ejecutivo", level=1)
    doc.add_paragraph(
        "Se desarrollo y desplego una aplicacion web para registrar ingredientes disponibles en casa "
        "y obtener recetas generadas o recomendadas a partir del inventario. El proyecto cumple el stack "
        "solicitado: Python con FastAPI, PostgreSQL, integracion configurable con OpenRouter como proveedor LLM, "
        "Docker Compose, VPS en AWS Lightsail, dominio propio y SSL."
    )
    add_key_table(doc)

    doc.add_heading("Que se construyo", level=1)
    add_bullets(
        doc,
        [
            "Backend modular en FastAPI con routers para autenticacion, ingredientes, recetas, recomendaciones y calificaciones.",
            "Base de datos PostgreSQL en produccion con modelos SQLAlchemy para usuarios, ingredientes, recetas y ratings.",
            "Interfaz web tipo dashboard con registro/login, inventario, preferencias culinarias, recomendaciones y recetas guardadas.",
            "Servicio LLM encapsulado en app/services/llm_service.py, preparado para OpenRouter y modelo openai/gpt-4o-mini.",
            "Motor local de recomendaciones con productos de canasta familiar, compatibilidad, faltantes y desglose de ingredientes.",
            "Documentacion automatica OpenAPI/Swagger disponible en /docs.",
        ],
    )

    doc.add_heading("Mejoras realizadas durante el desarrollo", level=1)
    add_bullets(
        doc,
        [
            "Se amplio el catalogo de productos comunes de canasta familiar para mejorar las recomendaciones.",
            "Se agrego desglose de ingredientes por receta: disponibles, faltantes, opcionales y basicos asumidos.",
            "Se enriquecio la experiencia visual con imagenes, favicon, tarjetas de recetas y un dashboard mas presentable.",
            "Se agregaron favoritos, calificaciones, busqueda, historial y copia de recetas al portapapeles.",
            "Se mejoro el PDF de entrega con arquitectura, diagrama ER, capturas y enlaces clicables.",
            "Se actualizaron README y documentos de arquitectura con URLs reales y datos del grupo.",
        ],
    )

    doc.add_heading("Despliegue y dominio", level=1)
    doc.add_paragraph(
        "La aplicacion se desplego en una VPS Ubuntu de AWS Lightsail con IP publica 54.236.36.56. "
        "Se uso Docker Compose para levantar tres servicios principales: app, db y caddy. Caddy funciona "
        "como proxy inverso y emite certificados SSL automaticos con Let's Encrypt."
    )
    add_numbered(
        doc,
        [
            "Se clono el repositorio en la VPS y se configuraron variables de entorno en .env.",
            "Se levanto PostgreSQL, FastAPI/Uvicorn y Caddy con Docker Compose.",
            "Se abrieron puertos 80 y 443 en Lightsail para HTTP/HTTPS.",
            "Primero se uso DuckDNS, pero algunas redes lo bloqueaban por ser dynamic DNS.",
            "Se uso sslip.io como respaldo tecnico mientras se resolvia el dominio propio.",
            "Finalmente se compro recetasruben.xyz en Namecheap y se configuraron registros A para @ y www hacia 54.236.36.56.",
            "Caddy emitio SSL de Let's Encrypt para recetasruben.xyz y www.recetasruben.xyz.",
        ],
    )

    doc.add_heading("Configuracion DNS final", level=1)
    dns_table = doc.add_table(rows=1, cols=4)
    dns_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    dns_table.style = "Table Grid"
    for idx, header in enumerate(["Tipo", "Host", "Valor", "TTL"]):
        set_cell_text(dns_table.rows[0].cells[idx], header, True)
        set_cell_shading(dns_table.rows[0].cells[idx], "F2F4F7")
    for row in [("A Record", "@", "54.236.36.56", "Automatic"), ("A Record", "www", "54.236.36.56", "Automatic")]:
        cells = dns_table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)

    doc.add_heading("Validaciones realizadas", level=1)
    add_bullets(
        doc,
        [
            "La aplicacion en produccion respondio HTTP 200 en https://www.recetasruben.xyz.",
            "Swagger respondio HTTP 200 en https://www.recetasruben.xyz/docs.",
            "El dominio sin www tambien respondio correctamente en https://recetasruben.xyz.",
            "El certificado SSL fue emitido por Let's Encrypt y validado por Caddy.",
            "La base de datos PostgreSQL quedo healthy en Docker Compose.",
            "Los tests locales del proyecto pasaron: 14 passed.",
            "Los cambios de PDF, README, Caddyfile y documentos fueron subidos a GitHub.",
        ],
    )

    doc.add_heading("Como se debe sustentar", level=1)
    add_bullets(
        doc,
        [
            "Explicar que el usuario registra ingredientes y el sistema recomienda o genera recetas segun el inventario.",
            "Mostrar el flujo: registro/login, agregar ingredientes, ver recomendaciones, generar receta y guardar/calificar.",
            "Mostrar Swagger para evidenciar endpoints REST y documentacion automatica.",
            "Explicar que OpenRouter esta encapsulado en un servicio interno y no se expone como chatbot libre.",
            "Explicar la arquitectura Docker: Caddy SSL -> FastAPI -> PostgreSQL, con OpenRouter como proveedor externo.",
            "Mostrar el PDF tecnico con arquitectura, diagrama ER y capturas.",
        ],
    )

    doc.add_heading("Checklist final para enviar", level=1)
    checklist = doc.add_table(rows=1, cols=3)
    checklist.alignment = WD_TABLE_ALIGNMENT.CENTER
    checklist.style = "Table Grid"
    for idx, header in enumerate(["Elemento", "Estado", "Evidencia"]):
        set_cell_text(checklist.rows[0].cells[idx], header, True)
        set_cell_shading(checklist.rows[0].cells[idx], "F2F4F7")
    for row in [
        ("Repositorio publico", "Listo", "GitHub con commits y documentos actualizados"),
        ("Aplicacion en produccion", "Listo", "https://www.recetasruben.xyz responde 200"),
        ("SSL", "Listo", "Certificado Let's Encrypt emitido por Caddy"),
        ("Swagger", "Listo", "https://www.recetasruben.xyz/docs responde 200"),
        ("PDF tecnico", "Listo", "Incluye arquitectura, ER, capturas y enlaces"),
        ("Documento explicativo", "Listo", "Este documento resume desarrollo y despliegue"),
    ]:
        cells = checklist.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, bold=(idx == 0))

    doc.add_heading("Evidencias tecnicas rapidas", level=1)
    add_bullets(
        doc,
        [
            "Docker Compose quedo con servicios app, db y caddy levantados.",
            "La base de datos quedo healthy en el contenedor de PostgreSQL.",
            "Caddy sirve HTTPS para recetasruben.xyz y www.recetasruben.xyz.",
            "El PDF tecnico y este documento quedaron dentro de la carpeta docs/.",
            "El repositorio remoto quedo actualizado en la rama main.",
        ],
    )

    doc.add_heading("Notas finales", level=1)
    doc.add_paragraph(
        "El repositorio no contiene credenciales reales. Las claves sensibles se manejan por variables de entorno. "
        "El dominio propio y SSL ya cumplen el requisito del parcial. La URL de respaldo solo se conserva como "
        "plan B tecnico, pero la URL principal para entrega es el dominio propio recetasruben.xyz."
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Generador de Recetas con Inventario - Entrega final").font.size = Pt(9)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
